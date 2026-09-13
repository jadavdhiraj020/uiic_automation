import os
import openpyxl
import pytest
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
from app.data.data_model import ClaimData
from app.data.excel_reader import _calculate_professional_fee
from app.data.folder_scanner import scan_folder, _extract_sheet_for_reinspection
from app.utils import load_doc_mapping, load_field_mapping, load_automation_defaults


@pytest.fixture(autouse=True)
def mock_excel_com_for_tests(monkeypatch):
    """Ensure tests in this module do not invoke real COM automation on Excel unless explicitly tested."""
    def fake_generate(excel_path, output_pdf_path, logs=None, *args, **kwargs):
        Path(output_pdf_path).write_bytes(b"%PDF-1.4 mock auto pdf")
        return True
    monkeypatch.setattr("app.data.printable_excel_service._generate_pdf_excel_com", fake_generate)



def test_survey_report_auto_copies_to_assessment_report(tmp_path):
    """
    Requirement 1:
    - User only configures/uploads Survey Report.
    - Assessment report setting is removed from doc_mapping claim_assessment_tab.
    - When Survey Report is detected, create a physical copy assessment_report.pdf,
      keep original as survey_report, and map both to assessment_files.
    """
    doc_map = load_doc_mapping(portal_id="uiic")
    assert "assessment_report" not in doc_map.get("claim_assessment_tab", {})
    assert "survey_report" in doc_map.get("claim_assessment_tab", {})

    # Create dummy survey report
    survey_file = tmp_path / "survey_report_final.pdf"
    survey_file.write_bytes(b"%PDF-1.4 dummy survey content")

    result = scan_folder(str(tmp_path), portal_id="uiic")

    # Verify both are mapped in assessment_files
    assert "survey_report" in result.assessment_files
    assert "assessment_report" in result.assessment_files
    assert os.path.basename(result.assessment_files["survey_report"]) == "survey_report_final.pdf"
    assert os.path.basename(result.assessment_files["assessment_report"]) == "assessment_report.pdf"

    # Verify physical file was created and contains identical bytes
    copied_path = Path(result.assessment_files["assessment_report"])
    assert copied_path.exists()
    assert copied_path.read_bytes() == b"%PDF-1.4 dummy survey content"
    assert str(copied_path) in result.generated_files


def test_vehicle_photographs_single_prefix_maps_to_4_slots(tmp_path):
    """
    Requirement 2:
    - doc_mapping has single 'Vehicle Photographs' key under claim_documents_tab.
    - Single photo_sheet.jpg is detected and duplicated into 4 files.
    - Mapped to the 4 angle slots required by portal automation.
    """
    doc_map = load_doc_mapping(portal_id="uiic")
    claim_tab = doc_map.get("claim_documents_tab", {})
    assert "Vehicle Photographs" in claim_tab
    assert "Vehicle Photograph (Front)" not in claim_tab
    assert "Vehicle Photograph(Rear)" not in claim_tab

    # Create single photo
    photo_file = tmp_path / "photo_sheet_damaged_car.jpg"
    photo_file.write_bytes(b"JPEG dummy image data")

    result = scan_folder(str(tmp_path), portal_id="uiic")

    expected_slots = [
        "Vehicle Photograph (Front)",
        "Vehicle Photograph(Rear)",
        "Vehicle Photograph (Left)",
        "Vehicle Photograph (Right)",
    ]

    for slot in expected_slots:
        assert slot in result.claim_doc_files
        slot_file = Path(result.claim_doc_files[slot])
        assert slot_file.exists()
        assert slot_file.read_bytes() == b"JPEG dummy image data"

    # Verify 4 physical copy files exist
    for i in range(1, 5):
        assert (tmp_path / f"vehicle_photo_{i}.jpg").exists()


def test_reinspection_3rd_sheet_extraction_and_manual_priority(tmp_path, monkeypatch):
    """
    Requirement 3:
    - Extracts 3rd sheet (index 2) from main Excel into reinspection.xlsx.
    - Prioritizes user-provided manual reinspection file if present.
    - Ensures reinspection.xlsx is excluded from main Excel candidates.
    """
    # Isolate from user AppData overrides
    from app.utils import doc_mapping_paths, read_json_file
    bundled = read_json_file(doc_mapping_paths(portal_id="uiic")["default"])
    monkeypatch.setattr("app.utils.load_doc_mapping", lambda portal_id=None: bundled)

    # 1. Create a 3-sheet workbook
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Summary"
    ws1["A1"] = "Main Data"

    ws2 = wb.create_sheet(title="Estimate")
    ws2["A1"] = "Estimate Details"

    ws3 = wb.create_sheet(title="Re-inspection")
    ws3["A1"] = "Reinspection Data Content"
    ws3["B2"] = 12345

    excel_path = tmp_path / "claim_main_data.xlsx"
    wb.save(str(excel_path))

    # Scan folder without user manual reinspection
    result = scan_folder(str(tmp_path), portal_id="uiic")

    assert result.excel_path == str(excel_path)
    assert "reinspection_report" in result.assessment_files
    reinspection_path = Path(result.assessment_files["reinspection_report"])
    assert reinspection_path.name == "reinspection.xlsx"
    assert reinspection_path.exists()

    # Verify extracted sheet content
    wb_extracted = openpyxl.load_workbook(str(reinspection_path))
    assert len(wb_extracted.sheetnames) == 1
    assert wb_extracted.active["A1"].value == "Reinspection Data Content"
    assert wb_extracted.active["B2"].value == 12345

    # 2. Test manual reinspection file priority
    manual_pdf = tmp_path / "manual_reinspection_report.pdf"
    manual_pdf.write_bytes(b"%PDF manual reinspection")

    result2 = scan_folder(str(tmp_path), portal_id="uiic")
    assert result2.excel_path == str(excel_path)
    assert "reinspection_report" in result2.assessment_files
    assert Path(result2.assessment_files["reinspection_report"]).name == "manual_reinspection_report.pdf"


def test_professional_fee_calculation_and_field_mapping():
    """
    Requirement 4:
    - field_mapping.json defines survey_fee and reinspection_fee.
    - ClaimData model has survey_fee, reinspection_fee, and professional_fee.
    - professional_fee = survey_fee + reinspection_fee via safe Decimal arithmetic.
    """
    fm = load_field_mapping(portal_id="uiic")
    assert "survey_fee" in fm
    assert "reinspection_fee" in fm
    assert "professional_fee" not in fm

    # Test fee addition
    claim = ClaimData()
    claim.survey_fee = "1500.50"
    claim.reinspection_fee = "500.25"
    claim.professional_fee = _calculate_professional_fee(claim.survey_fee, claim.reinspection_fee)
    assert claim.professional_fee == "2000.75"

    # Test edge cases (missing, comma formatting, invalid)
    assert _calculate_professional_fee("2,500.00", "500") == "3000"
    assert _calculate_professional_fee("1200", "") == "1200"
    assert _calculate_professional_fee(None, "450.50") == "450.50"
    assert _calculate_professional_fee("N/A", "invalid") == "0"
    assert _calculate_professional_fee("", "") == "0"

    # Verify preview fields include all 3 fees
    preview_fields = dict((f[0], f[1]) for f in claim.all_fields_for_preview())
    assert "Survey Fee (₹)" in preview_fields
    assert "Reinspection Fee (₹)" in preview_fields
    assert "Prof. Fee (Total ₹)" in preview_fields
    assert preview_fields["Survey Fee (₹)"] == "1500.50"
    assert preview_fields["Reinspection Fee (₹)"] == "500.25"
    assert preview_fields["Prof. Fee (Total ₹)"] == "2000.75"


def test_email_id_regex_first_and_fallback(tmp_path):
    """
    Test Email ID discovery:
    1. Regex-first finds .com email (e.g. cell I4 or anywhere on Sheet1) without requiring labels.
    2. Fallback to standard label-based search if no .com regex match is present.
    """
    from app.data.excel_reader import extract_claim_data

    # Scenario 1: Email at cell I4 (Row 4, Column I -> index r=3, c=8) without any label
    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.title = "Sheet1"
    ws1["A1"] = "Claim Number"
    ws1["B1"] = "CLM-998877"
    # Row 4, Column I is I4
    ws1["I4"] = "surveyor_office@gmail.com"
    file1 = tmp_path / "email_test_i4.xlsx"
    wb1.save(str(file1))

    claim1 = extract_claim_data(str(file1), portal_id="uiic")
    assert claim1.email_id == "surveyor_office@gmail.com"
    assert "I4 (Sheet1)" in claim1._excel_coords.get("email_id", "")

    # Scenario 2: Embedded email inside text cell without explicit "email" label
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "Sheet1"
    ws2["C2"] = "Head Office Contact: support.survey@insurance-audit.com (Direct)"
    file2 = tmp_path / "email_test_regex_embedded.xlsx"
    wb2.save(str(file2))

    claim2 = extract_claim_data(str(file2), portal_id="uiic")
    assert claim2.email_id == "support.survey@insurance-audit.com"

    # Scenario 3: No .com email, but standard label "email" exists with non-.com email (fallback)
    wb3 = openpyxl.Workbook()
    ws3 = wb3.active
    ws3.title = "Sheet1"
    ws3["A5"] = "email"
    ws3["B5"] = "surveyor@portal.gov.in"
    file3 = tmp_path / "email_test_fallback_label.xlsx"
    wb3.save(str(file3))

    claim3 = extract_claim_data(str(file3), portal_id="uiic")
    assert claim3.email_id == "surveyor@portal.gov.in"

    # Scenario 4: No email anywhere in workbook
    wb4 = openpyxl.Workbook()
    ws4 = wb4.active
    ws4.title = "Sheet1"
    ws4["A1"] = "Some Header"
    file4 = tmp_path / "email_test_empty.xlsx"
    wb4.save(str(file4))

    claim4 = extract_claim_data(str(file4), portal_id="uiic")
    assert claim4.email_id == ""


def test_expected_completion_date_calculation_and_preview():
    """
    Requirement (Point 14):
    - expected_completion_date is dependent ONLY and ONLY on date_of_survey.
    - If date_of_survey is present: expected_completion_date = date_of_survey + 1 calendar month.
    - Handles leap years and month-end clamping (e.g. 31/01/2024 -> 29/02/2024, 31/01/2026 -> 28/02/2026).
    - If date_of_survey is not present / empty: both date_of_survey and expected_completion_date remain blank.
    """
    from app.data.data_model import _add_one_calendar_month, ClaimData

    # 1. Calendar math tests
    # User's actual date from UI screenshot: 27/04/2023 -> 27/05/2023
    assert _add_one_calendar_month("27/04/2023") == "27/05/2023"
    assert _add_one_calendar_month("15/05/2026") == "15/06/2026"
    # Non-leap year February month-end (28 days)
    assert _add_one_calendar_month("31/01/2026") == "28/02/2026"
    assert _add_one_calendar_month("29/01/2026") == "28/02/2026"
    # Leap year February month-end (29 days in 2024)
    assert _add_one_calendar_month("31/01/2024") == "29/02/2024"
    assert _add_one_calendar_month("29/01/2024") == "29/02/2024"
    # 31st to 30-day month
    assert _add_one_calendar_month("31/03/2026") == "30/04/2026"
    assert _add_one_calendar_month("31/08/2026") == "30/09/2026"
    assert _add_one_calendar_month("31/10/2026") == "30/11/2026"
    # Year rollover (December -> January)
    assert _add_one_calendar_month("15/12/2026") == "15/01/2027"
    assert _add_one_calendar_month("31/12/2026") == "31/01/2027"
    # Empty or invalid inputs
    assert _add_one_calendar_month("") == ""
    assert _add_one_calendar_month(None) == ""
    assert _add_one_calendar_month("invalid-date") == ""

    # 2. ClaimData integration: When date_of_survey is present
    claim = ClaimData()
    claim.date_of_survey = "27/04/2023"
    claim._excel_coords["date_of_survey"] = "R65C6 (Survey report)"
    claim.final_report_date = "01/09/2026"  # Should NOT affect expected_completion_date
    claim.calculate_derived_fields()

    assert claim.expected_completion_date == "27/05/2023"
    assert claim._excel_coords.get("expected_completion_date") == "R65C6 (Survey report)"

    # Verify preview dictionary for UIIC
    preview_dict = dict((item[0], item[1]) for item in claim.all_fields_for_preview())
    assert preview_dict["Date of Survey"] == "27/04/2023"
    assert preview_dict["Expected Compl. Date"] == "27/05/2023"

    # 3. ClaimData integration: When date_of_survey is NOT present (both blank)
    claim_empty = ClaimData()
    claim_empty.date_of_survey = ""
    claim_empty.final_report_date = "01/09/2026"  # Even if report date exists, must stay blank!
    claim_empty.calculate_derived_fields()

    assert claim_empty.date_of_survey == ""
    assert claim_empty.expected_completion_date == ""
    assert "expected_completion_date" not in claim_empty._excel_coords

    preview_empty = dict((item[0], item[1]) for item in claim_empty.all_fields_for_preview())
    assert preview_empty["Date of Survey"] == ""
    assert preview_empty["Expected Compl. Date"] == ""


def test_professional_fee_cross_portal_isolation(tmp_path):
    """
    Ensure professional_fee calculation does NOT break New India or OIC portals:
    - New India: extracts professional_fee directly from Excel, must NOT be overwritten to 0.
    - OIC: extracts professional_fee directly from Excel, must NOT be overwritten to 0.
    - UIIC: derives professional_fee = survey_fee + reinspection_fee.
    """
    from app.data.excel_reader import extract_claim_data

    # 1. New India workbook with extracted professional_fee
    wb_nia = openpyxl.Workbook()
    ws_nia = wb_nia.active
    ws_nia.title = "Sheet1"
    ws_nia["A1"] = "professional fee"
    ws_nia["B1"] = "3500"
    file_nia = tmp_path / "nia_prof_fee.xlsx"
    wb_nia.save(str(file_nia))

    claim_nia = extract_claim_data(str(file_nia), portal_id="newindia")
    assert claim_nia.professional_fee == "3500", "New India professional_fee must not be overwritten to 0!"

    # 2. OIC workbook with extracted professional_fee
    wb_oic = openpyxl.Workbook()
    ws_oic = wb_oic.active
    ws_oic.title = "Sheet1"
    ws_oic["A1"] = "professional fee"
    ws_oic["B1"] = "4200"
    file_oic = tmp_path / "oic_prof_fee.xlsx"
    wb_oic.save(str(file_oic))

    claim_oic = extract_claim_data(str(file_oic), portal_id="oic")
    assert claim_oic.professional_fee == "4200", "OIC professional_fee must not be overwritten to 0!"

    # 3. UIIC workbook with survey_fee and reinspection_fee
    wb_uiic = openpyxl.Workbook()
    ws_uiic = wb_uiic.active
    ws_uiic.title = "Sheet1"
    ws_uiic["A1"] = "Survey Fee"
    ws_uiic["B1"] = "2000"
    ws_uiic["A2"] = "Reinspection Fee"
    ws_uiic["B2"] = "750"
    file_uiic = tmp_path / "uiic_prof_fee.xlsx"
    wb_uiic.save(str(file_uiic))

    claim_uiic = extract_claim_data(str(file_uiic), portal_id="uiic")
    assert claim_uiic.survey_fee == "2000"
    assert claim_uiic.reinspection_fee == "750"
    assert claim_uiic.professional_fee == "2750", "UIIC professional_fee must equal survey_fee + reinspection_fee"


def test_survey_time_extraction_and_fallback(tmp_path, monkeypatch):
    """
    Test survey time extraction and fallback:
    1. Robust _parse_time_string for 12h and 24h formats.
    2. Extraction when time is in the date cell.
    3. Extraction when time is in an adjacent cell.
    4. Fallback to 11 HH and 00 MM when time is missing in Excel.
    5. Blank when date_of_survey is missing.
    """
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    from app.data.excel_reader import _parse_time_string, extract_claim_data

    # 1. Parsing tests
    assert _parse_time_string("11:30 AM") == ("11", "30")
    assert _parse_time_string("11.30am") == ("11", "30")
    assert _parse_time_string("02:45 PM") == ("14", "45")
    assert _parse_time_string("2 PM") == ("14", "00")
    assert _parse_time_string("14:30") == ("14", "30")
    assert _parse_time_string("2023-04-27 15:45:00") == ("15", "45")
    assert _parse_time_string("") is None
    assert _parse_time_string("no time here") is None

    # 2. Time in date cell
    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.title = "Sheet1"
    ws1["A1"] = "Date and Time of Survey"
    ws1["E1"] = "27/04/2023 15:30"
    f1 = tmp_path / "time_in_cell.xlsx"
    wb1.save(str(f1))

    c1 = extract_claim_data(str(f1), portal_id="uiic")
    assert c1.time_hh == "15"
    assert c1.time_mm == "30"

    # 3. Dedicated "Time" field on same row
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "Sheet1"
    ws2["A1"] = "Date and Time of Survey"
    ws2["E1"] = "27/04/2023"
    ws2["F1"] = "Time"
    ws2["G1"] = "02:15 PM"
    f2 = tmp_path / "time_separate.xlsx"
    wb2.save(str(f2))

    c2 = extract_claim_data(str(f2), portal_id="uiic")
    assert c2.date_of_survey == "27/04/2023"
    assert c2.time_hh == "14"
    assert c2.time_mm == "15"

    # 3b. Dedicated "Time of Survey" field on next row
    wb2b = openpyxl.Workbook()
    ws2b = wb2b.active
    ws2b.title = "Sheet1"
    ws2b["A1"] = "Date and Time of Survey"
    ws2b["E1"] = "27/04/2023"
    ws2b["A2"] = "Time of Survey"
    ws2b["E2"] = "15:30"
    f2b = tmp_path / "time_row2.xlsx"
    wb2b.save(str(f2b))

    c2b = extract_claim_data(str(f2b), portal_id="uiic")
    assert c2b.date_of_survey == "27/04/2023"
    assert c2b.time_hh == "15"
    assert c2b.time_mm == "30"

    # 4. Missing time in Excel -> Fallback to 11:00
    wb3 = openpyxl.Workbook()
    ws3 = wb3.active
    ws3.title = "Sheet1"
    ws3["A1"] = "Date and Time of Survey"
    ws3["E1"] = "27/04/2023"
    f3 = tmp_path / "time_missing.xlsx"
    wb3.save(str(f3))

    c3 = extract_claim_data(str(f3), portal_id="uiic")
    assert c3.date_of_survey == "27/04/2023"
    assert c3.time_hh == "11"
    assert c3.time_mm == "00"
    preview3 = dict((item[0], item[1]) for item in c3.all_fields_for_preview())
    assert preview3["Time of Survey"] == "11:00"

    # 5. Missing date_of_survey -> both date and time remain blank
    wb4 = openpyxl.Workbook()
    ws4 = wb4.active
    ws4.title = "Sheet1"
    ws4["A1"] = "Claim Number"
    ws4["B1"] = "CLM-123"
    f4 = tmp_path / "no_date.xlsx"
    wb4.save(str(f4))

    c4 = extract_claim_data(str(f4), portal_id="uiic")
    assert c4.date_of_survey == ""
    assert c4.time_hh == ""
    assert c4.time_mm == ""
    preview4 = dict((item[0], item[1]) for item in c4.all_fields_for_preview())
    assert preview4["Time of Survey"] == ""


def test_gst_summary_parts_and_labour_removed_from_config_and_preview():
    """
    Verifies that gst_summary_parts and gst_summary_labour are completely removed:
    1. Not present in UIIC field_mapping.json.
    2. Not present in UI preview table for UIIC workspace.
    """
    from app.utils import load_field_mapping
    from app.data.data_model import ClaimData

    # 1. Field mapping configuration check
    mapping = load_field_mapping(portal_id="uiic")
    assert "gst_summary_parts" not in mapping
    assert "gst_summary_labour" not in mapping

    # 2. UI preview table check
    claim = ClaimData(portal_id="uiic")
    preview_labels = [item[0] for item in claim.all_fields_for_preview()]
    assert "GST Summary Parts (₹)" not in preview_labels
    assert "GST Summary Labour (₹)" not in preview_labels



def test_survey_report_auto_generated_from_excel_when_missing(tmp_path, monkeypatch):
    """
    Test auto-generation of survey_report.pdf from Excel when missing:
    1. Folder has only data.xlsx (no manual survey PDF).
    2. scan_folder generates survey_report.pdf from Sheet 1 and clones it to assessment_report.pdf.
    3. Both are mapped to assessment_files without timestamp suffixes.
    4. Re-scanning overwrites cleanly without creating duplicate timestamped files.
    """
    from unittest.mock import patch

    # Isolate from user AppData overrides
    from app.utils import doc_mapping_paths, read_json_file
    bundled = read_json_file(doc_mapping_paths(portal_id="uiic")["default"])
    monkeypatch.setattr("app.utils.load_doc_mapping", lambda portal_id=None: bundled)

    # 1. Create a dummy Excel workbook
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Survey report"
    ws1["A1"] = "MOTOR FINAL SURVEY REPORT"
    ws1["B2"] = "Claim Data"

    excel_path = tmp_path / "claim_data.xlsx"
    wb.save(str(excel_path))
    wb.close()

    # Mock PDF generator to write a mock PDF
    def fake_pdf_gen(excel_path, pdf_path, logs):
        with open(pdf_path, "wb") as f:
            f.write(b"%PDF-1.4 mock survey report content")
        return True

    with patch("app.data.printable_excel_service._generate_pdf_excel_com", side_effect=fake_pdf_gen):
        # First scan
        result = scan_folder(str(tmp_path), portal_id="uiic")

        # Verify survey_report and assessment_report are both generated and mapped
        assert "survey_report" in result.assessment_files
        assert "assessment_report" in result.assessment_files

        survey_path = Path(result.assessment_files["survey_report"])
        assessment_path = Path(result.assessment_files["assessment_report"])

        assert survey_path.name == "survey_report.pdf"
        assert assessment_path.name == "assessment_report.pdf"
        assert survey_path.exists()
        assert assessment_path.exists()

        # Both files contain the generated content
        assert survey_path.read_bytes() == b"%PDF-1.4 mock survey report content"
        assert assessment_path.read_bytes() == b"%PDF-1.4 mock survey report content"

        # Second scan (re-scan)
        result2 = scan_folder(str(tmp_path), portal_id="uiic")

        # Verify no duplicate timestamped files were created in the folder
        files = os.listdir(tmp_path)
        pdf_files = [f for f in files if f.endswith(".pdf")]

        assert sorted(pdf_files) == ["assessment_report.pdf", "survey_report.pdf"]
        assert len([f for f in files if "Printable_Assessment" in f]) == 0


def test_invoice_pdf_extraction_disabled_and_excel_retained(tmp_path):
    """
    Verify that invoice PDF text/OCR extraction is disabled, so Excel values
    are authoritative and never overwritten by invoice PDF files.
    """
    from app.ui.services.claim_folder_service import ClaimFolderService
    from app.data.folder_scanner import FolderScanResult

    service = ClaimFolderService(portal_id="uiic")

    # Mock an invoice PDF file in tmp_path
    invoice_pdf = tmp_path / "workshop_invoice.pdf"
    invoice_pdf.write_bytes(b"%PDF-1.4 dummy invoice text with Invoice No: PDF-INV-999")

    claim = ClaimData()
    claim.invoice_no = "EXCEL-INV-100"
    claim.invoice_date = "12/03/2026"
    claim.workshop_invoice_no = "WS-INV-555"
    claim.workshop_invoice_date = "10/03/2026"

    scan_result = FolderScanResult()
    scan_result.folder_path = str(tmp_path)
    scan_result.assessment_files = {"invoice": str(invoice_pdf)}

    logs = []
    service._extract_pdf_invoice_data(scan_result, claim, logs)

    # Values must remain their own distinct Excel values with ZERO mixing
    assert claim.invoice_no == "EXCEL-INV-100"
    assert claim.workshop_invoice_no == "WS-INV-555"
    assert claim.workshop_invoice_date == "10/03/2026"
    assert getattr(claim, "_pending_invoice_ocr", False) is False


def test_excel_invoice_and_workshop_invoice_independent_validation():
    """
    Verify that invoice_no and workshop_invoice_no are validated independently
    without any cross-fallback or mixing.
    """
    claim = ClaimData()
    claim.invoice_no = "EXCEL-INV-200"
    claim.invoice_date = "15/04/2026"
    claim.workshop_invoice_no = "WS-300"
    claim.workshop_invoice_date = "14/04/2026"
    claim.assessment_files = {"invoice": "mock_invoice.pdf"}

    # Validation should not complain about invoice missing when both are present
    errors, warnings = claim._validate_uiic()
    assert not any("Workshop Invoice No" in w for w in warnings)
    assert not any("Workshop Invoice Date" in w for w in warnings)

    # If workshop_invoice_no is empty, the workshop warning MUST appear even if invoice_no is set
    partial_claim = ClaimData()
    partial_claim.invoice_no = "EXCEL-INV-200"
    partial_claim.invoice_date = "15/04/2026"
    partial_claim.assessment_files = {"invoice": "mock_invoice.pdf"}
    _, partial_warnings = partial_claim._validate_uiic()
    assert "Workshop Invoice No not found in Excel" in partial_warnings
    assert "Workshop Invoice Date not found in Excel" in partial_warnings


def test_background_generation_skips_pdf_when_already_exists(tmp_path):
    """
    Verify Issue 3 fix: If survey_report.pdf was already generated during the scan,
    process_printable_output skips redundant PDF render via printable_skip_pdf setting.
    """
    from app.data.printable_excel_service import process_printable_output

    excel_path = tmp_path / "data.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Data"
    wb.save(str(excel_path))
    wb.close()

    logs = []
    settings = {
        "printable_output_excel_name": "Printable_Assessment.xlsx",
        "printable_output_pdf_name": "survey_report.pdf",
        "printable_skip_pdf": True,
    }

    process_printable_output(
        source_excel_path=str(excel_path),
        output_folder=str(tmp_path),
        settings=settings,
        logs=logs,
    )
    assert any("Printable PDF skipped" in line for line in logs)


def test_uiic_clean_number_suffix_extraction():
    """
    Verify that invoice_no, workshop_invoice_no, and final_report_no
    are cleaned to their trailing number suffix for UIIC without affecting other portals:
    - 'INV-2026-0042' -> '0042' (preserves leading zeros)
    - 'JDB/2026-27/PORTAL/8744' -> '8744'
    - 'ABSD/2026-27/PORTAL/12434' -> '12434'
    - '8744' -> '8744'
    """
    from app.data.data_model import ClaimData, _extract_number_suffix

    # 1. Helper function checks
    assert _extract_number_suffix("INV-2026-0042") == "0042"
    assert _extract_number_suffix("JDB/2026-27/PORTAL/8744") == "8744"
    assert _extract_number_suffix("ABSD/2026-27/PORTAL/12434") == "12434"
    assert _extract_number_suffix("8744") == "8744"
    assert _extract_number_suffix("") == ""

    # 2. UIIC portal post-processing
    c_uiic = ClaimData(portal_id="uiic")
    c_uiic.invoice_no = "INV-2026-0042"
    c_uiic.final_report_no = "JDB/2026-27/PORTAL/8744"
    c_uiic.workshop_invoice_no = "INV-2026-0042"
    c_uiic.calculate_derived_fields()

    assert c_uiic.invoice_no == "0042"
    assert c_uiic.final_report_no == "8744"
    assert c_uiic.workshop_invoice_no == "0042"

    preview_uiic = dict((item[0], item[1]) for item in c_uiic.all_fields_for_preview())
    assert preview_uiic["Invoice No"] == "0042"
    assert preview_uiic["Report No"] == "8744"

    # 3. Isolation: Other portals (e.g. New India, OIC) are untouched
    c_nia = ClaimData(portal_id="newindia")
    c_nia.invoice_no = "INV-2026-0042"
    c_nia.final_report_no = "JDB/2026-27/PORTAL/8744"
    c_nia.calculate_derived_fields()

    assert c_nia.invoice_no == "INV-2026-0042"
    assert c_nia.final_report_no == "JDB/2026-27/PORTAL/8744"


def test_uiic_settings_pdf_mapping_toggles():
    """
    Verify that in SettingsPage:
    1. In UIIC mode: pdf_fields_container is hidden, pdf_uiic_info_container is visible.
    2. In non-UIIC mode (e.g. New India): pdf_fields_container is visible, info container is hidden.
    """
    import sys
    from PyQt6.QtWidgets import QApplication
    from app.ui.components.settings_page import SettingsPage

    app = QApplication.instance()
    if not app:
        app = QApplication(sys.argv)

    page = SettingsPage()
    # Default is uiic: fields container hidden, info container visible
    assert page._portal_id == "uiic"
    assert page.pdf_fields_container.isHidden() is True
    assert page.pdf_uiic_info_container.isHidden() is False

    # Switch to newindia: fields container visible, info container hidden
    page.set_portal("newindia")
    assert page._portal_id == "newindia"
    assert page.pdf_fields_container.isHidden() is False
    assert page.pdf_uiic_info_container.isHidden() is True

    # Switch back to uiic
    page.set_portal("uiic")
    assert page._portal_id == "uiic"
    assert page.pdf_fields_container.isHidden() is True
    assert page.pdf_uiic_info_container.isHidden() is False


def test_generate_pdf_excel_com_sets_pagesetup_print_area_for_column_range(monkeypatch):
    """
    Reason 2 verification: When print_mode="column_range", Excel COM directly
    configures ws.PageSetup.PrintArea = f"${start_c}:${end_c}", ws.PageSetup.Zoom = False,
    and FitToPagesWide = 1 to guarantee the print area is honored by the COM print engine.
    """
    monkeypatch.undo()
    from unittest.mock import MagicMock, patch
    from app.data.printable_excel_service import _generate_pdf_excel_com

    mock_excel = MagicMock()
    mock_wb = MagicMock()
    mock_ws = MagicMock()

    mock_excel.Workbooks.Open.return_value = mock_wb
    mock_wb.Worksheets.return_value = mock_ws
    mock_dispatch = MagicMock(return_value=mock_excel)

    mock_win32 = MagicMock()
    mock_win32.client = MagicMock()
    mock_win32.client.DispatchEx = mock_dispatch
    mock_pythoncom = MagicMock()

    with patch.dict("sys.modules", {
        "win32com": mock_win32,
        "win32com.client": mock_win32.client,
        "pythoncom": mock_pythoncom,
    }):
        with patch("app.data.printable_excel_service._find_pdf_printer", return_value="Microsoft Print to PDF"):
            logs = []
            res = _generate_pdf_excel_com(
                "dummy.xlsx",
                "dummy.pdf",
                logs,
                print_mode="column_range",
                col_range=("A", "B"),
            )

    assert res is True
    assert mock_ws.PageSetup.PrintArea == "$A:$B"
    assert mock_ws.PageSetup.Zoom is False
    assert mock_ws.PageSetup.FitToPagesWide == 1
    assert mock_ws.PageSetup.FitToPagesTall is False
    mock_ws.ExportAsFixedFormat.assert_called_once()


def test_generate_pdf_excel_com_sets_pagesetup_zoom_for_scale_percentage(monkeypatch):
    """
    Reason 2 verification: When print_mode="scale_percentage", Excel COM directly
    clears ws.PageSetup.PrintArea and sets ws.PageSetup.Zoom to the integer scale.
    """
    monkeypatch.undo()
    from unittest.mock import MagicMock, patch
    from app.data.printable_excel_service import _generate_pdf_excel_com

    mock_excel = MagicMock()
    mock_wb = MagicMock()
    mock_ws = MagicMock()

    mock_excel.Workbooks.Open.return_value = mock_wb
    mock_wb.Worksheets.return_value = mock_ws
    mock_dispatch = MagicMock(return_value=mock_excel)

    mock_win32 = MagicMock()
    mock_win32.client = MagicMock()
    mock_win32.client.DispatchEx = mock_dispatch
    mock_pythoncom = MagicMock()

    with patch.dict("sys.modules", {
        "win32com": mock_win32,
        "win32com.client": mock_win32.client,
        "pythoncom": mock_pythoncom,
    }):
        with patch("app.data.printable_excel_service._find_pdf_printer", return_value="Microsoft Print to PDF"):
            logs = []
            res = _generate_pdf_excel_com(
                "dummy.xlsx",
                "dummy.pdf",
                logs,
                print_mode="scale_percentage",
                scale=75,
            )

    assert res is True
    assert mock_ws.PageSetup.PrintArea == ""
    assert mock_ws.PageSetup.Zoom == 75
    mock_ws.ExportAsFixedFormat.assert_called_once()


def test_uiic_scan_folder_freshly_regenerates_and_overwrites_survey_report(tmp_path):
    """
    Reason 1 verification: Re-scanning freshly re-generates and overwrites
    survey_report.pdf and syncs to assessment_report.pdf instead of skipping
    when the file already exists on disk.
    """
    import openpyxl
    from unittest.mock import patch
    from app.data.folder_scanner import scan_folder

    # Create dummy Excel matching uiic excel_keywords ("main1")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Data"
    excel_path = tmp_path / "main1.xlsx"
    wb.save(str(excel_path))
    wb.close()

    # Pre-existing stale survey_report.pdf
    survey_path = tmp_path / "survey_report.pdf"
    with open(survey_path, "wb") as f:
        f.write(b"%PDF-1.4 OLD_REPORT")

    # Mock generator that writes NEW_REPORT
    def fake_render(excel_in, pdf_out, logs, **kwargs):
        with open(pdf_out, "wb") as f:
            f.write(b"%PDF-1.4 NEW_REPORT")
        return True

    with patch("app.data.printable_excel_service._generate_pdf_excel_com", side_effect=fake_render):
        result = scan_folder(str(tmp_path), portal_id="uiic")

    # Must be mapped
    assert result.assessment_files.get("survey_report") == str(survey_path)
    assessment_path = tmp_path / "assessment_report.pdf"
    assert result.assessment_files.get("assessment_report") == str(assessment_path)

    # Must be overwritten with NEW_REPORT
    with open(survey_path, "rb") as f:
        assert f.read() == b"%PDF-1.4 NEW_REPORT"

    with open(assessment_path, "rb") as f:
        assert f.read() == b"%PDF-1.4 NEW_REPORT"



