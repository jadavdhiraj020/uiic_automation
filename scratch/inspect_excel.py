import openpyxl
import os

path = r"c:\Users\jadav\Coding\Automation\UIIC\uiic_automation\AAA\dummy.xlsx"
if os.path.exists(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    print("Sheets:", wb.sheetnames)
    for name in wb.sheetnames:
        sheet = wb[name]
        print(f"\n--- Sheet: {name} ---")
        for r_idx, row in enumerate(sheet.iter_rows(values_only=True)):
            row_str = " | ".join(str(cell) if cell is not None else "" for cell in row)
            if "invoice" in row_str.lower() or "number" in row_str.lower() or "date" in row_str.lower():
                print(f"Row {r_idx+1}: {row_str[:200]}")
else:
    print("File not found")
