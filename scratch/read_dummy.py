import sys
import openpyxl

# Force stdout to use utf-8 to prevent UnicodeEncodeError on Windows
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r"c:\Users\jadav\Coding\Automation\UIIC\uiic_automation\AAA\dummy.xlsx", data_only=True)
sheet = wb.active
print("Active Sheet:", sheet.title)

for r in range(40, min(100, sheet.max_row + 1)):
    row_vals = [sheet.cell(r, c).value for c in range(1, min(15, sheet.max_column + 1))]
    # Print row if there is any value
    if any(v is not None for v in row_vals):
        print(f"Row {r:02d}:", [str(v) if v is not None else "" for v in row_vals])
