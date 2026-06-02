"""
oic_assessment_generator.py — Generate & read back OIC-format assessment Excel.

Website 3 (OIC Portal) uses this module to:
  1. Extract Spare Parts + Labour from the source Excel (shared extractor)
  2. Write an `auto_oic_assessment.xlsx` with the 8 OIC-portal-mirror columns
  3. Write a companion `auto_oic_assessment_audit.txt`
  4. Re-read the generated Excel to provide row dicts for portal filling

The generated Excel preserves the ORIGINAL source Excel row order (not grouped
by material type).  Portal filling groups by Item Type internally.

Column layout (mirrors what gets typed into the OIC portal):
  A: SlNo
  B: Item Type      (Glass / Plastic / Metallic Parts / Labour Charge)
  C: Item Sub Type   (part name / labour description)
  D: Side Description (config default, e.g. "FRONT")
  E: Item Amount     (billed amount)
  F: IGST Rate       (config default, e.g. "18%")
  G: Estimated Amount
  H: HSN Code
"""

import logging
import os
from dataclasses import dataclass
from typing import Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from app.data.excel_parts_extractor import (
    extract_parts_and_labour,
    coerce_hsn_code,
)

logger = logging.getLogger(__name__)

_OUTPUT_FILENAME = "auto_oic_assessment.xlsx"


@dataclass
class OicAssessmentGenerationResult:
    output_path: Optional[str]
    status: str
    generated_files: List[str]
    message: str = ""

# ── Material → OIC Portal "Item Type" text ────────────────────────────────────
_MATERIAL_TO_ITEM_TYPE: Dict[str, str] = {
    "glass":   "Glass",
    "plastic": "Plastic",
    "metal":   "Metallic Parts",
}


# ══════════════════════════════════════════════════════════════════════════════
# INTERNAL HELPERS
# ══════════════════════════════════════════════════════════════════════════════

def _next_available_path(output_folder: str, filename: str = _OUTPUT_FILENAME) -> str:
    """Return a non-conflicting file path in *output_folder*."""
    base, ext = os.path.splitext(filename)
    candidate = os.path.join(output_folder, filename)
    idx = 1
    while os.path.exists(candidate):
        candidate = os.path.join(output_folder, f"{base}_{idx}{ext}")
        idx += 1
    return candidate


def _audit_path_for(output_path: str) -> str:
    root, _ext = os.path.splitext(output_path)
    return f"{root}_audit.txt"


def _write_audit_file(audit_path: str, lines: List[str]) -> None:
    with open(audit_path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines).rstrip() + "\n")


def _safe_amount_str(value) -> str:
    """Convert a numeric amount to a clean string (1500.0 → '1500')."""
    try:
        num = float(value)
        if num == int(num) and not (num != num):  # guard NaN
            return str(int(num))
        return str(num)
    except (ValueError, TypeError, OverflowError):
        return str(value)


# ══════════════════════════════════════════════════════════════════════════════
# ROW TRANSFORMATION  (extractor rows → OIC 8-column rows)
# ══════════════════════════════════════════════════════════════════════════════

def _transform_rows(
    parts_rows: List[dict],
    labour_rows: List[dict],
    defaults: dict,
    parts_hsn: str,
    labour_hsn: str,
) -> List[dict]:
    """Transform extracted rows into OIC 8-column format.

    Preserves the original source Excel order:
      - parts_rows come first (in the order the extractor found them)
      - labour_rows come after (in the order the extractor found them)

    Each output dict has keys matching the 8-column schema:
      item_type, item_sub_type, side_description,
      item_amount, igst_rate, estimated_amount, hsn_code
    """
    side_desc = defaults.get("item_side_description", "FRONT")
    igst_rate = defaults.get("item_igst_rate", "18%")
    labour_est_amount = defaults.get("labour_estimated_amount", "100000")

    oic_rows: List[dict] = []

    # ── Spare Parts (original order) ──────────────────────────────────────────
    for row in parts_rows:
        material = row.get("material_type", "metal")
        item_type = _MATERIAL_TO_ITEM_TYPE.get(material, "Metallic Parts")

        oic_rows.append({
            "item_type": item_type,
            "item_sub_type": row["part_name"],
            "side_description": side_desc,
            "item_amount": _safe_amount_str(row["billed_amount"]),
            "igst_rate": igst_rate,
            "estimated_amount": _safe_amount_str(row["estimated_amount"]),
            "hsn_code": parts_hsn,
        })

    # ── Labour Charges (original order) ───────────────────────────────────────
    for row in labour_rows:
        oic_rows.append({
            "item_type": "Labour Charge",
            "item_sub_type": row["part_name"],
            "side_description": side_desc,
            "item_amount": _safe_amount_str(row["billed_amount"]),
            "igst_rate": igst_rate,
            "estimated_amount": labour_est_amount,
            "hsn_code": labour_hsn,
        })

    return oic_rows


# ══════════════════════════════════════════════════════════════════════════════
# EXCEL GENERATION
# ══════════════════════════════════════════════════════════════════════════════

_HEADERS = [
    "SlNo",
    "Item Type",
    "Item Sub Type",
    "Side Description",
    "Item Amount",
    "IGST Rate",
    "Estimated Amount",
    "HSN Code",
]

_HEADER_KEY_MAP = [
    "item_type",
    "item_sub_type",
    "side_description",
    "item_amount",
    "igst_rate",
    "estimated_amount",
    "hsn_code",
]

_COL_WIDTHS = {
    "A": 8, "B": 20, "C": 35, "D": 18,
    "E": 16, "F": 12, "G": 20, "H": 14,
}


def _create_oic_workbook(oic_rows: List[dict]):
    """Create an in-memory workbook with the OIC 8-column format."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "OIC_Assessment_Data"
    ws.views.sheetView[0].showGridLines = True

    # ── Header styling ────────────────────────────────────────────────────────
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, header_text in enumerate(_HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # ── Column widths ─────────────────────────────────────────────────────────
    for col_letter, width in _COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    # ── Data rows ─────────────────────────────────────────────────────────────
    for idx, row_data in enumerate(oic_rows, start=1):
        r = idx + 1  # row 1 is header
        ws.cell(row=r, column=1, value=idx)  # SlNo
        for col_offset, key in enumerate(_HEADER_KEY_MAP, 2):
            ws.cell(row=r, column=col_offset, value=row_data[key])

    return wb


# ══════════════════════════════════════════════════════════════════════════════
# PUBLIC API: GENERATE
# ══════════════════════════════════════════════════════════════════════════════

def generate_oic_assessment(
    source_excel_path: str,
    output_folder: str,
    defaults: dict,
) -> Optional[str]:
    """Generate ``auto_oic_assessment.xlsx`` from the source Excel.

    Args:
        source_excel_path: Path to the main data Excel file.
        output_folder:     Directory where the generated file will be saved.
        defaults:          OIC automation defaults dict (from config).

    Returns:
        Full path to the generated file, or ``None`` if no data was extracted.

    Raises:
        FileNotFoundError: If *source_excel_path* does not exist.
    """
    if not os.path.exists(source_excel_path):
        raise FileNotFoundError(f"Source Excel not found: {source_excel_path}")

    parts_hsn = coerce_hsn_code(defaults.get("item_hsn_code"), "8512")
    labour_hsn = coerce_hsn_code(defaults.get("labour_hsn_code"), "8729")
    boundary_kw = defaults.get("table_boundary_keyword", "sub total").strip() or "sub total"

    logger.info("OIC assessment generator: loading source Excel %s", source_excel_path)

    # ── Phase 1: Extract ──────────────────────────────────────────────────────
    parts_rows, labour_rows, extractor_audit = extract_parts_and_labour(
        source_excel_path, parts_hsn=parts_hsn, labour_hsn=labour_hsn,
        boundary_keyword=boundary_kw,
    )
    logger.info(
        "OIC assessment generator: extracted %d parts + %d labour rows",
        len(parts_rows), len(labour_rows),
    )

    audit_lines: List[str] = [
        "OIC Assessment Generation Audit",
        f"Source Excel: {source_excel_path}",
        f"Parts HSN: {parts_hsn}",
        f"Labour HSN: {labour_hsn}",
        "",
    ]
    audit_lines.extend(extractor_audit)

    # ── Guard: no data ────────────────────────────────────────────────────────
    if not parts_rows and not labour_rows:
        logger.warning("OIC assessment generator: no data found — skipping generation")
        # Use a fixed filename so we do not consume a numbered slot that the
        # actual generator would later claim, avoiding phantom path conflicts.
        audit_path = os.path.join(output_folder, "oic_assessment_skipped_audit.txt")
        audit_lines.extend(["", "Result", "No parts or labour data found. Excel generation skipped."])
        _write_audit_file(audit_path, audit_lines)
        return None

    # ── Phase 2: Transform to OIC 8-column format ─────────────────────────────
    oic_rows = _transform_rows(parts_rows, labour_rows, defaults, str(parts_hsn), str(labour_hsn))
    logger.info("OIC assessment generator: transformed %d OIC rows", len(oic_rows))

    # ── Phase 3: Generate Excel ───────────────────────────────────────────────
    wb = _create_oic_workbook(oic_rows)

    output_path = _next_available_path(output_folder)
    audit_path = _audit_path_for(output_path)
    _max_save_attempts = 3

    try:
        for _attempt in range(1, _max_save_attempts + 1):
            try:
                wb.save(output_path)
                break
            except PermissionError:
                if _attempt == _max_save_attempts:
                    logger.error(
                        "OIC assessment generator: all %d save attempts failed (last path: %s)",
                        _max_save_attempts, output_path,
                    )
                    audit_lines.extend([
                        "", "Result",
                        f"FAILED: Could not save after {_max_save_attempts} attempts.",
                        f"Last attempted path: {output_path}",
                    ])
                    _write_audit_file(audit_path, audit_lines)
                    raise
                logger.warning(
                    "OIC assessment generator: path locked on attempt %d, retrying: %s",
                    _attempt, output_path,
                )
                output_path = _next_available_path(output_folder)
                audit_path = _audit_path_for(output_path)
    finally:
        wb.close()

    # ── Audit summary ─────────────────────────────────────────────────────────
    audit_lines.extend([
        "",
        "── OIC Row Details ──",
    ])
    for idx, row in enumerate(oic_rows, 1):
        audit_lines.append(
            f"  {idx}. [{row['item_type']}] {row['item_sub_type']} | "
            f"amt={row['item_amount']} est={row['estimated_amount']} hsn={row['hsn_code']}"
        )

    parts_count = sum(1 for r in oic_rows if r["item_type"] != "Labour Charge")
    labour_count = sum(1 for r in oic_rows if r["item_type"] == "Labour Charge")
    audit_lines.extend([
        "",
        "── Summary ──",
        f"Output Excel: {output_path}",
        f"Audit File: {audit_path}",
        f"Parts rows: {parts_count}",
        f"Labour rows: {labour_count}",
        f"Total rows: {len(oic_rows)}",
    ])
    _write_audit_file(audit_path, audit_lines)

    logger.info(
        "OIC assessment generator: SUCCESS — %d rows (%d parts + %d labour) → %s",
        len(oic_rows), parts_count, labour_count, output_path,
    )
    return output_path


def generate_oic_assessment_result(
    source_excel_path: str,
    output_folder: str,
    defaults: dict,
) -> OicAssessmentGenerationResult:
    output_path = generate_oic_assessment(source_excel_path, output_folder, defaults)
    if output_path:
        return OicAssessmentGenerationResult(
            output_path=output_path,
            status="generated",
            generated_files=[output_path, _audit_path_for(output_path)],
            message="OIC assessment generated.",
        )

    skipped_audit = os.path.join(output_folder, "oic_assessment_skipped_audit.txt")
    generated_files = [skipped_audit] if os.path.exists(skipped_audit) else []
    return OicAssessmentGenerationResult(
        output_path=None,
        status="skipped_no_data",
        generated_files=generated_files,
        message="Assessment generation skipped: no parts/labour data found.",
    )


# ══════════════════════════════════════════════════════════════════════════════
# PUBLIC API: READ BACK
# ══════════════════════════════════════════════════════════════════════════════

def read_oic_assessment(oic_excel_path: str) -> List[dict]:
    """Read a generated OIC assessment Excel back into row dicts.

    Each returned dict has keys:
        item_type, item_sub_type, side_description,
        item_amount, igst_rate, estimated_amount, hsn_code

    Args:
        oic_excel_path: Path to ``auto_oic_assessment.xlsx``.

    Returns:
        List of row dicts (empty list if the file has no data rows).

    Raises:
        FileNotFoundError: If *oic_excel_path* does not exist.
    """
    if not os.path.exists(oic_excel_path):
        raise FileNotFoundError(f"OIC assessment Excel not found: {oic_excel_path}")

    logger.info("OIC assessment reader: loading %s", oic_excel_path)
    wb = openpyxl.load_workbook(oic_excel_path, data_only=True)
    ws = wb.active

    rows: List[dict] = []

    for r in range(2, ws.max_row + 1):
        # Column B = Item Type (col 2)
        item_type = ws.cell(row=r, column=2).value
        if item_type is None or str(item_type).strip() == "":
            continue  # skip empty rows

        igst_raw = ws.cell(row=r, column=6).value
        if igst_raw is None or str(igst_raw).strip() == "":
            logger.warning(
                "OIC assessment reader: IGST Rate cell is blank at row %d in '%s' — defaulting to empty string",
                r, oic_excel_path,
            )
        rows.append({
            "item_type":        str(ws.cell(row=r, column=2).value or "").strip(),
            "item_sub_type":    str(ws.cell(row=r, column=3).value or "").strip(),
            "side_description": str(ws.cell(row=r, column=4).value or "").strip(),
            "item_amount":      str(ws.cell(row=r, column=5).value or "0").strip(),
            "igst_rate":        str(igst_raw or "").strip(),
            "estimated_amount": str(ws.cell(row=r, column=7).value or "0").strip(),
            "hsn_code":         str(ws.cell(row=r, column=8).value or "").strip(),
        })

    wb.close()
    logger.info("OIC assessment reader: read %d rows from %s", len(rows), oic_excel_path)
    return rows
